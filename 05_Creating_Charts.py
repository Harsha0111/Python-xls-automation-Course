from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference

# Create a workbook
wb = Workbook()

# Create the first sheet for the bar chart data
ws_bar = wb.active
ws_bar.title = "Bar Chart Data"

# Add data for the bar chart
ws_bar.append(["Category", "Value"])
ws_bar.append(["A", 10])
ws_bar.append(["B", 20])
ws_bar.append(["C", 30])

# Create a bar chart and add it to the bar chart sheet
bar_chart = BarChart()
bar_data = Reference(ws_bar, min_col=2, min_row=2, max_row=4)
bar_chart.add_data(bar_data, titles_from_data=True)
bar_chart.title = "Bar Chart Example"
ws_bar.add_chart(bar_chart, "E5")

# Create a new sheet for the pie chart data
ws_pie = wb.create_sheet(title="Pie Chart Data")

# Add data for the pie chart
ws_pie.append(["Category", "Percentage"])
ws_pie.append(["X", 40])
ws_pie.append(["Y", 60])

# Create a pie chart and add it to the pie chart sheet
pie_chart = PieChart()
pie_data = Reference(ws_pie, min_col=2, min_row=2, max_row=3)
pie_chart.add_data(pie_data, titles_from_data=True)
pie_chart.title = "Pie Chart Example"
ws_pie.add_chart(pie_chart, "E5")

# Save workbook with charts in separate sheets
wb.save("output/05_Creating_Charts.xlsx")
print("Workbook with charts saved as 'output/05_Creating_Charts.xlsx'")
